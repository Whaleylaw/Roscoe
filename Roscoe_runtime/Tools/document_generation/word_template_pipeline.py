#!/usr/bin/env python3
"""
Word Template Pipeline

High-fidelity Word template filling using docxtpl (Jinja2-style placeholders).
Supports .docx template rendering and .doc/.docx → PDF conversion via LibreOffice.

This module provides the agent with tools to:
1. Fill attorney-approved Word templates with case data
2. Convert .doc files to .docx (legacy template support)
3. Export .docx to PDF with layout preservation

Usage:
    from word_template_pipeline import (
        render_docx_template,
        convert_docx_to_pdf,
        convert_doc_to_docx,
        fill_and_export_template
    )
    
    # Fill template and export to PDF
    result = fill_and_export_template(
        template_path="/path/to/template.docx",
        output_path="/path/to/output.docx",
        context={"client_name": "John Smith", "date_of_loss": "01/15/2024"},
        export_pdf=True
    )
"""

import os
import shutil
import subprocess
import sys
from datetime import datetime, date
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

try:
    from docxtpl import DocxTemplate, InlineImage
    DOCXTPL_AVAILABLE = True
except ImportError:
    DOCXTPL_AVAILABLE = False

try:
    from docx import Document
    from docx.shared import Inches
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    PYTHON_DOCX_AVAILABLE = False


class TemplateError(Exception):
    """Base exception for template operations."""
    pass


class TemplateNotFoundError(TemplateError):
    """Template file not found."""
    pass


class LibreOfficeNotFoundError(TemplateError):
    """LibreOffice not installed or not in PATH."""
    pass


class ConversionError(TemplateError):
    """Document conversion failed."""
    pass


class PlaceholderError(TemplateError):
    """Placeholder resolution failed."""
    pass


def _find_libreoffice() -> Optional[str]:
    """
    Find LibreOffice executable path.
    
    Returns:
        Path to soffice executable, or None if not found.
    """
    # Common locations
    candidates = [
        "soffice",  # In PATH (Linux, macOS with brew)
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",  # macOS default
        "/usr/bin/soffice",  # Linux
        "/usr/bin/libreoffice",  # Linux alternative
        "/snap/bin/libreoffice",  # Ubuntu snap
    ]
    
    for candidate in candidates:
        if shutil.which(candidate):
            return candidate
        if os.path.isfile(candidate) and os.access(candidate, os.X_OK):
            return candidate
    
    return None


def check_dependencies() -> Dict[str, bool]:
    """
    Check availability of required dependencies.
    
    Returns:
        Dict with dependency name -> available boolean
    """
    return {
        "docxtpl": DOCXTPL_AVAILABLE,
        "python-docx": PYTHON_DOCX_AVAILABLE,
        "libreoffice": _find_libreoffice() is not None,
    }


def sanitize_context(context: Dict[str, Any]) -> Dict[str, Any]:
    """
    Sanitize context values for safe template rendering.
    
    Converts:
    - None -> empty string
    - datetime/date objects -> formatted strings
    - Decimal -> float
    - Lists/dicts -> recursively sanitized
    - Other types -> string representation
    
    Args:
        context: Raw context dictionary
    
    Returns:
        Sanitized context dictionary safe for Jinja2 rendering
    """
    def _sanitize_value(value: Any) -> Any:
        if value is None:
            return ""
        elif isinstance(value, str):
            return value
        elif isinstance(value, bool):
            return value
        elif isinstance(value, (int, float)):
            return value
        elif isinstance(value, Decimal):
            return float(value)
        elif isinstance(value, datetime):
            return value.strftime("%B %d, %Y")
        elif isinstance(value, date):
            return value.strftime("%B %d, %Y")
        elif isinstance(value, dict):
            return {k: _sanitize_value(v) for k, v in value.items()}
        elif isinstance(value, (list, tuple)):
            return [_sanitize_value(item) for item in value]
        else:
            return str(value)
    
    return {key: _sanitize_value(val) for key, val in context.items()}


def render_docx_template(
    template_path: Union[str, Path],
    context: Dict[str, Any],
    output_path: Union[str, Path],
    autoescape: bool = True,
    sanitize: bool = True
) -> Dict[str, Any]:
    """
    Render a DOCX template with Jinja2-style placeholders.
    
    Placeholders in the template should use Jinja2 syntax:
    - {{ variable_name }} for simple substitution
    - {% for item in items %} ... {% endfor %} for loops
    - {% if condition %} ... {% endif %} for conditionals
    
    Args:
        template_path: Path to the .docx template file
        context: Dictionary of values to fill placeholders
        output_path: Path for the output .docx file
        autoescape: Enable XML auto-escaping (recommended for safety)
        sanitize: Sanitize context values before rendering
    
    Returns:
        Dict with status, output_path, and any warnings
    
    Raises:
        TemplateNotFoundError: If template file doesn't exist
        TemplateError: If rendering fails
    """
    if not DOCXTPL_AVAILABLE:
        raise TemplateError(
            "docxtpl not installed. Run: pip install docxtpl"
        )
    
    template_path = Path(template_path)
    output_path = Path(output_path)
    
    if not template_path.exists():
        raise TemplateNotFoundError(f"Template not found: {template_path}")
    
    result = {
        "status": "success",
        "output_path": str(output_path),
        "warnings": [],
        "placeholders_found": []
    }
    
    try:
        # Load template
        doc = DocxTemplate(str(template_path))
        
        # Get placeholders from template (for reporting)
        try:
            placeholders = doc.get_undeclared_template_variables()
            result["placeholders_found"] = list(placeholders)
        except Exception:
            # Some templates may fail variable extraction
            pass
        
        # Sanitize context if requested
        if sanitize:
            context = sanitize_context(context)
        
        # Check for missing placeholders
        if result["placeholders_found"]:
            missing = [p for p in result["placeholders_found"] if p not in context]
            if missing:
                result["warnings"].append(
                    f"Missing context keys for placeholders: {missing}"
                )
        
        # Render template
        doc.render(context, autoescape=autoescape)
        
        # Ensure output directory exists
        output_path.parent.mkdir(parents=True, exist_ok=True)
        
        # Save rendered document
        doc.save(str(output_path))
        
        return result
        
    except Exception as e:
        raise TemplateError(f"Template rendering failed: {e}") from e


def convert_doc_to_docx(
    input_path: Union[str, Path],
    output_dir: Optional[Union[str, Path]] = None,
    timeout: int = 120
) -> Dict[str, Any]:
    """
    Convert a .doc file to .docx using LibreOffice.
    
    Args:
        input_path: Path to the .doc file
        output_dir: Directory for output (defaults to same as input)
        timeout: Conversion timeout in seconds
    
    Returns:
        Dict with status, output_path, and any errors
    
    Raises:
        LibreOfficeNotFoundError: If LibreOffice is not installed
        ConversionError: If conversion fails
    """
    soffice = _find_libreoffice()
    if not soffice:
        raise LibreOfficeNotFoundError(
            "LibreOffice not found. Install with:\n"
            "  macOS: brew install --cask libreoffice\n"
            "  Ubuntu: sudo apt-get install libreoffice"
        )
    
    input_path = Path(input_path)
    if not input_path.exists():
        raise TemplateNotFoundError(f"Input file not found: {input_path}")
    
    if output_dir is None:
        output_dir = input_path.parent
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    
    result = {
        "status": "success",
        "output_path": None,
        "error": None
    }
    
    try:
        cmd = [
            soffice,
            "--headless",
            "--convert-to", "docx",
            "--outdir", str(output_dir),
            str(input_path)
        ]
        
        process = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=timeout
        )
        
        # LibreOffice uses input filename with new extension
        expected_output = output_dir / f"{input_path.stem}.docx"
        
        if expected_output.exists():
            result["output_path"] = str(expected_output)
            return result
        else:
            error_msg = process.stderr or "LibreOffice did not produce output file"
            raise ConversionError(f".doc → .docx conversion failed: {error_msg}")
            
    except subprocess.TimeoutExpired:
        raise ConversionError(f"Conversion timed out after {timeout} seconds")
    except subprocess.SubprocessError as e:
        raise ConversionError(f"Conversion process failed: {e}")


def convert_docx_to_pdf(
    input_path: Union[str, Path],
    output_dir: Optional[Union[str, Path]] = None,
    timeout: int = 120
) -> Dict[str, Any]:
    """
    Convert a .docx file to PDF using LibreOffice.
    
    This method preserves Word formatting and styles better than
    pure Python libraries like reportlab.
    
    Args:
        input_path: Path to the .docx file
        output_dir: Directory for PDF output (defaults to same as input)
        timeout: Conversion timeout in seconds
    
    Returns:
        Dict with status, output_path, and any errors
    
    Raises:
        LibreOfficeNotFoundError: If LibreOffice is not installed
        ConversionError: If conversion fails
    """
    soffice = _find_libreoffice()
    if not soffice:
        raise LibreOfficeNotFoundError(
            "LibreOffice not found. Install with:\n"
            "  macOS: brew install --cask libreoffice\n"
            "  Ubuntu: sudo apt-get install libreoffice"
        )
    
    input_path = Path(input_path)
    if not input_path.exists():
        raise TemplateNotFoundError(f"Input file not found: {input_path}")
    
    if output_dir is None:
        output_dir = input_path.parent
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    
    result = {
        "status": "success",
        "output_path": None,
        "error": None
    }
    
    try:
        cmd = [
            soffice,
            "--headless",
            "--convert-to", "pdf",
            "--outdir", str(output_dir),
            str(input_path)
        ]
        
        process = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=timeout
        )
        
        # LibreOffice uses input filename with .pdf extension
        expected_output = output_dir / f"{input_path.stem}.pdf"
        
        if expected_output.exists():
            result["output_path"] = str(expected_output)
            return result
        else:
            error_msg = process.stderr or "LibreOffice did not produce PDF"
            raise ConversionError(f".docx → PDF conversion failed: {error_msg}")
            
    except subprocess.TimeoutExpired:
        raise ConversionError(f"Conversion timed out after {timeout} seconds")
    except subprocess.SubprocessError as e:
        raise ConversionError(f"Conversion process failed: {e}")


def fill_and_export_template(
    template_path: Union[str, Path],
    output_path: Union[str, Path],
    context: Dict[str, Any],
    export_pdf: bool = True,
    pdf_output_dir: Optional[Union[str, Path]] = None
) -> Dict[str, Any]:
    """
    Fill a Word template and optionally export to PDF.
    
    This is the main entry point for the template pipeline.
    
    Args:
        template_path: Path to the .docx template
        output_path: Path for the filled .docx output
        context: Dictionary of values for placeholders
        export_pdf: Whether to also generate a PDF
        pdf_output_dir: Directory for PDF (defaults to same as docx)
    
    Returns:
        Dict with:
            - status: "success" or "error"
            - docx_path: Path to generated .docx
            - pdf_path: Path to generated .pdf (if export_pdf=True)
            - warnings: List of non-fatal warnings
            - errors: List of error messages (if status="error")
    
    Example:
        result = fill_and_export_template(
            template_path="templates/medical_request.docx",
            output_path="output/John_Smith_Medical_Request.docx",
            context={
                "client_name": "John Smith",
                "provider_name": "ABC Medical",
                "date_of_loss": "01/15/2024",
                "today_date": "12/12/2025"
            },
            export_pdf=True
        )
    """
    result = {
        "status": "success",
        "docx_path": None,
        "pdf_path": None,
        "warnings": [],
        "errors": []
    }
    
    try:
        # Step 1: Render template
        render_result = render_docx_template(
            template_path=template_path,
            context=context,
            output_path=output_path
        )
        
        result["docx_path"] = render_result["output_path"]
        result["warnings"].extend(render_result.get("warnings", []))
        
        # Step 2: Convert to PDF if requested
        if export_pdf:
            try:
                pdf_result = convert_docx_to_pdf(
                    input_path=output_path,
                    output_dir=pdf_output_dir
                )
                result["pdf_path"] = pdf_result["output_path"]
            except LibreOfficeNotFoundError as e:
                result["warnings"].append(f"PDF export skipped: {e}")
            except ConversionError as e:
                result["warnings"].append(f"PDF export failed: {e}")
        
        return result
        
    except TemplateError as e:
        result["status"] = "error"
        result["errors"].append(str(e))
        return result
    except Exception as e:
        result["status"] = "error"
        result["errors"].append(f"Unexpected error: {e}")
        return result


def list_template_placeholders(template_path: Union[str, Path]) -> List[str]:
    """
    Extract placeholder variable names from a template.
    
    Args:
        template_path: Path to the .docx template
    
    Returns:
        List of placeholder names found in the template
    """
    if not DOCXTPL_AVAILABLE:
        raise TemplateError("docxtpl not installed")
    
    template_path = Path(template_path)
    if not template_path.exists():
        raise TemplateNotFoundError(f"Template not found: {template_path}")
    
    doc = DocxTemplate(str(template_path))
    try:
        return list(doc.get_undeclared_template_variables())
    except Exception:
        return []


# ============================================================================
# CLI Interface
# ============================================================================

def main():
    """Command-line interface for the word template pipeline."""
    import argparse
    import json
    
    parser = argparse.ArgumentParser(
        description="Word Template Pipeline - Fill templates and convert to PDF",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  # Check dependencies
  python word_template_pipeline.py --check-deps
  
  # List placeholders in a template
  python word_template_pipeline.py --list-placeholders template.docx
  
  # Fill template with JSON context
  python word_template_pipeline.py --template template.docx --output filled.docx \\
      --context '{"client_name": "John Smith", "date_of_loss": "01/15/2024"}'
  
  # Fill template and export to PDF
  python word_template_pipeline.py --template template.docx --output filled.docx \\
      --context-file context.json --pdf
  
  # Convert .doc to .docx
  python word_template_pipeline.py --convert-doc legacy.doc
  
  # Convert .docx to PDF only
  python word_template_pipeline.py --to-pdf document.docx
        """
    )
    
    parser.add_argument("--check-deps", action="store_true",
                        help="Check if dependencies are available")
    parser.add_argument("--list-placeholders", metavar="TEMPLATE",
                        help="List placeholders in a template")
    parser.add_argument("--template", "-t", metavar="PATH",
                        help="Path to template .docx file")
    parser.add_argument("--output", "-o", metavar="PATH",
                        help="Output path for filled document")
    parser.add_argument("--context", "-c", metavar="JSON",
                        help="Context as JSON string")
    parser.add_argument("--context-file", "-f", metavar="PATH",
                        help="Path to JSON file with context")
    parser.add_argument("--pdf", action="store_true",
                        help="Also export to PDF")
    parser.add_argument("--convert-doc", metavar="PATH",
                        help="Convert .doc file to .docx")
    parser.add_argument("--to-pdf", metavar="PATH",
                        help="Convert .docx to PDF")
    parser.add_argument("--outdir", metavar="DIR",
                        help="Output directory for conversions")
    
    args = parser.parse_args()
    
    # Check dependencies
    if args.check_deps:
        deps = check_dependencies()
        print(json.dumps(deps, indent=2))
        return 0 if all(deps.values()) else 1
    
    # List placeholders
    if args.list_placeholders:
        try:
            placeholders = list_template_placeholders(args.list_placeholders)
            print(json.dumps(placeholders, indent=2))
            return 0
        except Exception as e:
            print(f"Error: {e}", file=sys.stderr)
            return 1
    
    # Convert .doc to .docx
    if args.convert_doc:
        try:
            result = convert_doc_to_docx(args.convert_doc, args.outdir)
            print(json.dumps(result, indent=2))
            return 0 if result["status"] == "success" else 1
        except Exception as e:
            print(f"Error: {e}", file=sys.stderr)
            return 1
    
    # Convert .docx to PDF
    if args.to_pdf:
        try:
            result = convert_docx_to_pdf(args.to_pdf, args.outdir)
            print(json.dumps(result, indent=2))
            return 0 if result["status"] == "success" else 1
        except Exception as e:
            print(f"Error: {e}", file=sys.stderr)
            return 1
    
    # Fill template
    if args.template:
        if not args.output:
            print("Error: --output required with --template", file=sys.stderr)
            return 1
        
        # Get context
        context = {}
        if args.context:
            try:
                context = json.loads(args.context)
            except json.JSONDecodeError as e:
                print(f"Error parsing context JSON: {e}", file=sys.stderr)
                return 1
        elif args.context_file:
            try:
                with open(args.context_file, 'r', encoding='utf-8') as f:
                    context = json.load(f)
            except Exception as e:
                print(f"Error reading context file: {e}", file=sys.stderr)
                return 1
        
        try:
            result = fill_and_export_template(
                template_path=args.template,
                output_path=args.output,
                context=context,
                export_pdf=args.pdf
            )
            print(json.dumps(result, indent=2))
            return 0 if result["status"] == "success" else 1
        except Exception as e:
            print(f"Error: {e}", file=sys.stderr)
            return 1
    
    parser.print_help()
    return 0


if __name__ == "__main__":
    sys.exit(main())
